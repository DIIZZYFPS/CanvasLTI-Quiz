import io
import os
import sys
from collections import Counter

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

from app.utils.docx_reader import extract_docx_paragraphs, _letters
from app.utils.file_reader import read_file
from app.utils.parser import parse_quiz_text

_next_id = [100]


def _fresh_id():
    _next_id[0] += 1
    return _next_id[0]


def add_list(doc, levels, start=1, overrides=None, abstract_id=None):
    """Define a list in the document's numbering part and return its numId.

    `levels` is a list of (numFmt, lvlText) for ilvl 0, 1, 2...; pass `abstract_id`
    to create another list instance (w:num) over an existing definition.
    """
    numbering = doc.part.numbering_part.element
    if abstract_id is None:
        abstract_id = _fresh_id()
        lvls = "".join(
            f'<w:lvl w:ilvl="{i}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/></w:lvl>'
            for i, (fmt, text) in enumerate(levels)
        )
        abstract = parse_xml(
            f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
            f'<w:multiLevelType w:val="multilevel"/>{lvls}</w:abstractNum>'
        )
        # abstractNum elements must precede num elements in numbering.xml
        numbering.find(qn("w:num")).addprevious(abstract)

    num_id = _fresh_id()
    override_xml = "".join(
        f'<w:lvlOverride w:ilvl="{lvl}"><w:startOverride w:val="{val}"/></w:lvlOverride>'
        for lvl, val in (overrides or {}).items()
    )
    numbering.append(parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/>{override_xml}</w:num>'
    ))
    doc._test_abstract_ids = getattr(doc, "_test_abstract_ids", {})
    doc._test_abstract_ids[num_id] = abstract_id
    return num_id


def item(doc_or_cell, text, num_id, ilvl=0):
    p = doc_or_cell.add_paragraph(text)
    numpr = p._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_numId().val = num_id
    numpr.get_or_add_ilvl().val = ilvl
    return p


def lines(doc):
    return extract_docx_paragraphs(doc)


# --- label resolution --------------------------------------------------------

def test_style_based_numbered_list():
    """'List Number' carries its numbering on the style, not the paragraph."""
    doc = Document()
    for t in ("alpha", "beta", "gamma"):
        doc.add_paragraph(t, style="List Number")
    assert lines(doc) == ["1. alpha", "2. beta", "3. gamma"]


def test_quiz_with_auto_numbered_questions_and_auto_lettered_options():
    """The real-world shape: a multilevel list where questions are 1. 2. and the
    options under each are A. B. C., none of which are in the paragraph text."""
    doc = Document()
    n = add_list(doc, [("decimal", "%1."), ("upperLetter", "%2.")])
    item(doc, "What is 2+2?", n, 0)
    for t in ("three", "four", "five"):
        item(doc, t, n, 1)
    doc.add_paragraph("Answer: B")
    item(doc, "Capital of France?", n, 0)
    for t in ("Paris", "Lyon"):
        item(doc, t, n, 1)
    doc.add_paragraph("Answer: A")

    assert lines(doc) == [
        "1. What is 2+2?", "A) three", "B) four", "C) five", "Answer: B",
        "2. Capital of France?", "A) Paris", "B) Lyon", "Answer: A",   # options restart at A
    ]


def test_such_a_document_parses_into_real_questions():
    doc = Document()
    n = add_list(doc, [("decimal", "%1."), ("upperLetter", "%2.")])
    item(doc, "What is 2+2?", n, 0)
    for t in ("three", "four", "five"):
        item(doc, t, n, 1)
    doc.add_paragraph("Answer: B")
    item(doc, "Capital of France?", n, 0)
    for t in ("Paris", "Lyon"):
        item(doc, t, n, 1)
    doc.add_paragraph("Answer: A")
    buf = io.BytesIO()
    doc.save(buf)

    class Upload:
        filename = "quiz.docx"
        content_type = ""
        def read(self):
            return buf.getvalue()

    questions = parse_quiz_text(read_file(Upload()))
    assert [q["type"] for q in questions] == ["multiple_choice_question"] * 2
    assert [q["question_text"] for q in questions] == ["What is 2+2?", "Capital of France?"]
    for q, expected in zip(questions, ("four", "Paris")):
        correct = next(a for a in q["answers"] if a["id"] == q["correct_answer_id"])
        assert correct["text"] == expected
    assert [len(q["answers"]) for q in questions] == [3, 2]


def test_numbering_continues_across_unnumbered_paragraphs_in_the_same_list():
    doc = Document()
    n = add_list(doc, [("decimal", "%1.")])
    item(doc, "one", n)
    doc.add_paragraph("a note in between")
    item(doc, "two", n)
    assert lines(doc) == ["1. one", "a note in between", "2. two"]


def test_label_punctuation_is_normalised_to_what_the_parser_reads():
    """Letters always become 'A)' (the documented, routable option form) and
    numbers 'N.' or 'N)'; '(a)' and a bare 'A' are not forms the parser reads."""
    doc = Document()
    for fmt, text, expected in [
        ("lowerLetter", "(%1)", "a)"),
        ("upperLetter", "%1)", "A)"),
        ("upperLetter", "%1.", "A)"),
        ("upperLetter", "%1", "A)"),
        ("decimal", "%1.", "1."),
        ("decimal", "%1)", "1)"),
        ("decimal", "(%1)", "1)"),
        ("decimal", "%1", "1."),
    ]:
        n = add_list(doc, [(fmt, text)])
        item(doc, "x", n)
        assert lines(doc)[-1] == f"{expected} x", (fmt, text)


def test_start_value_is_respected():
    doc = Document()
    n = add_list(doc, [("decimal", "%1.")], start=5)
    item(doc, "x", n); item(doc, "y", n)
    assert lines(doc) == ["5. x", "6. y"]


def test_restart_numbering_override():
    doc = Document()
    first = add_list(doc, [("upperLetter", "%1.")])
    item(doc, "a1", first); item(doc, "a2", first)
    restarted = add_list(doc, None, abstract_id=doc._test_abstract_ids[first], overrides={0: 1})
    item(doc, "b1", restarted); item(doc, "b2", restarted)
    assert lines(doc) == ["A) a1", "B) a2", "A) b1", "B) b2"]


def test_letters_roll_over_after_z():
    assert [_letters(n, True) for n in (1, 26, 27, 52, 53)] == ["A", "Z", "AA", "AZ", "BA"]
    assert _letters(28, False) == "ab"


def test_bullets_and_roman_numerals_are_left_as_plain_text():
    doc = Document()
    doc.add_paragraph("bullet text", style="List Bullet")
    roman = add_list(doc, [("upperRoman", "%1.")])
    item(doc, "roman text", roman)
    assert lines(doc) == ["bullet text", "roman text"]


def test_numid_zero_removes_numbering_from_a_styled_paragraph():
    doc = Document()
    p = doc.add_paragraph("not actually numbered", style="List Number")
    p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    assert lines(doc) == ["not actually numbered"]


def test_empty_numbered_paragraph_still_counts():
    doc = Document()
    n = add_list(doc, [("upperLetter", "%1.")])
    item(doc, "x", n); item(doc, "", n); item(doc, "z", n)
    assert lines(doc) == ["A) x", "", "C) z"]


def test_document_without_lists_is_unchanged():
    doc = Document()
    for t in ("What is 2+2?", "A) 3", "B) 4", "Answer: B"):
        doc.add_paragraph(t)
    assert lines(doc) == ["What is 2+2?", "A) 3", "B) 4", "Answer: B"]


# --- tables ------------------------------------------------------------------

def test_table_text_is_included_in_reading_order():
    """Regression: document.paragraphs skips tables, so their text vanished silently."""
    doc = Document()
    doc.add_paragraph("Before")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"r{r}c{c}"
    doc.add_paragraph("After")
    assert lines(doc) == ["Before", "r0c0", "r0c1", "r1c0", "r1c1", "After"]


def test_merged_cells_are_read_once():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "merged"
    assert lines(doc) == ["merged"]


def test_numbered_list_inside_a_table_cell():
    doc = Document()
    n = add_list(doc, [("upperLetter", "%1.")])
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "Pick one"
    item(cell, "red", n); item(cell, "blue", n)
    assert lines(doc) == ["Pick one", "A) red", "B) blue"]


# --- compatibility -----------------------------------------------------------

def test_the_repos_sample_docx_parses_cleanly():
    """The sample used to report 1 error: its own '[Short Answer] ... Answer: Mars'
    example (also in the README) was misread as a fill-in-blank by the old bracket
    handling. It must now parse with no errors."""
    path = os.path.join(os.path.dirname(__file__), '..', 'Tests', 'QTI Test Case.docx')

    class Upload:
        filename = "QTI Test Case.docx"
        content_type = ""
        def read(self):
            with open(path, "rb") as f:
                return f.read()

    questions = parse_quiz_text(read_file(Upload()))
    assert len(questions) == 10
    assert Counter(q["type"] for q in questions) == {
        "multiple_choice_question": 2, "essay_question": 2, "true_false_question": 2,
        "short_answer_question": 4,
    }
